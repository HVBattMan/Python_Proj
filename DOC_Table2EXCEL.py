import docx
from docx import Document #导入库
import numpy as np
import pandas as pd

path = 'F:\\OneDrive\\SW Update\\Python_Script\\CN_NEV_Data\\22.docx'
document = Document(path)
tables = document.tables
# table = tables[0 ]#获取文件中的第一个表格
#获取变量名
k =1
#利用Pandas包中的ExcelWriter()方法增加一个公共句柄，在写入新的数据之时保留原来写入的数据，等到把所有的数据都写进去之后关闭这个句柄。
writer = pd.ExcelWriter('F:\\OneDrive\\SW Update\\Python_Script\\CN_NEV_Data\\22.xlsx')
for table in tables:
    variable_name = []
    for i in range(len(table.columns)):
        variable_name.append(table.cell(0, i).text)
    #获取表内数据
    data = []
    for i in range (1,len(table.rows)):
        for j in range (len(table.columns)):
            data.append(table.cell(i,j).text)

    #list to 1D data
    arr1 = np.array(data)
    #2d array
    arr2 = arr1.reshape(len(table.rows)-1,len(table.columns))
    #2d array to 2d sheet
    df = pd.DataFrame(arr2)
    df.columns = variable_name
    #output data
    sheet = 'Sheet ' +'%d'%(k)
    print (sheet)
    print (df)
    df.to_excel(writer, sheet_name= sheet)
    k = k+1
writer.close()